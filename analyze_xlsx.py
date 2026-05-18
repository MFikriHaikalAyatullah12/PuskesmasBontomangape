import openpyxl
import os

file_path = '/workspaces/PuskesmasBontomangape/FEBRUARI.xlsx'
if not os.path.exists(file_path):
    print(f"File {file_path} not found.")
    exit(1)

wb = openpyxl.load_parent_workbook = openpyxl.load_workbook(file_path, data_only=True)
print(f"Sheet names: {wb.sheetnames}")

for sheet_name in wb.sheetnames:
    print(f"\n--- Sheet: {sheet_name} ---")
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(max_row=15, values_only=True))
    
    for i, row in enumerate(rows):
        print(f"Row {i+1}: {row}")
    
    # Simple heuristic for header: first row with multiple non-empty values
    header_row = -1
    for i, row in enumerate(rows):
        non_empty = [str(val) for val in row if val is not None]
        if len(non_empty) > 3: # Heuristic
             header_row = i + 1
             break
    print(f"Likely header row index: {header_row}")
